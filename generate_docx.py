import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    val: 'single', 'double', 'dashed', etc.
    color: '00d2ff', '01244e', etc.
    sz: '4', '8', '12', '24'
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>'
            tcBorders.append(parse_xml(tag))
        else:
            tag = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def create_proposal_docx(output_path):
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles & Fonts
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(45, 55, 72)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    # Colors
    NAVY = RGBColor(1, 36, 78)       # #01244e
    BLUE = RGBColor(0, 100, 148)     # #006494
    CYAN = RGBColor(0, 160, 200)     # #00a0c8
    MUTED = RGBColor(100, 116, 139)  # #64748b

    # 1. Header & Title Block
    p_badge = doc.add_paragraph()
    r_badge = p_badge.add_run("PROPUESTA ESTRATÉGICA & DOSSIER MULTIMEDIA 2026")
    r_badge.font.size = Pt(8.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = CYAN
    p_badge.paragraph_format.space_after = Pt(2)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Radio Puerto Quidico 105.1 FM / 91.3 FM")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Transformación Digital: De Emisora Analógica a Estación Multimedia Comunitaria 24/7")
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = BLUE
    r_sub.font.bold = True
    p_sub.paragraph_format.space_after = Pt(6)

    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("Comuna de Tirúa, Región del Biobío • Borde Costero e Isla Mocha • Presentado a la Dirección General")
    r_meta.font.size = Pt(9)
    r_meta.font.italic = True
    r_meta.font.color.rgb = MUTED
    p_meta.paragraph_format.space_after = Pt(14)

    # Divider Line
    p_div = doc.add_paragraph()
    p_div_pPr = p_div._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="00A0C8"/></w:pBdr>')
    p_div_pPr.append(pBdr)
    p_div.paragraph_format.space_after = Pt(12)

    # 2. Resumen Ejecutivo
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Resumen Ejecutivo: ¿Qué se le entrega a la Radio?")
    r1.font.size = Pt(14)
    r1.font.color.rgb = NAVY
    r1.font.bold = True

    p = doc.add_paragraph(
        "Esta plataforma no es una simple página web informativa de presentación; es una estación interactiva completa "
        "diseñada a la medida de la realidad territorial, pesquera y cultural de Caleta Quidico y la Comuna de Tirúa. "
        "Resuelve las tres mayores limitantes históricas de la radiodifusión FM en la zona costera:"
    )

    bullets = [
        ("Fin a las zonas de sombra y lejanía: ", "La radio ahora llega a embarcaciones en faena marítima, hogares sin antena tradicional y a la diáspora familiar y lafkenche en Cañete, Concepción, Santiago o el extranjero."),
        ("Monetización medible y formal: ", "Los comerciantes ya no auspician solo por simpatía; la web registra cada clic y genera comprobantes de WhatsApp automáticos con la cantidad exacta de clientes que consultaron sus negocios."),
        ("Liderazgo en seguridad y servicio público: ", "Convierte a la radio en la voz indiscutible ante marejadas, cortes de camino y temporales, en coordinación directa con Capitanía de Puerto, Carabineros y SENAPRED.")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        r_bt = bp.add_run(b_title)
        r_bt.bold = True
        r_bt.font.color.rgb = BLUE
        bp.add_run(b_desc)

    # Callout Box
    table_callout = doc.add_table(rows=1, cols=1)
    table_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_c = table_callout.rows[0].cells[0]
    cell_c.width = Inches(6.7)
    set_cell_background(cell_c, "EDF8FD")
    set_cell_border(cell_c, left={'val': 'single', 'sz': '24', 'color': '00A0C8'})
    set_cell_margins(cell_c, top=140, bottom=140, left=180, right=180)
    
    cp = cell_c.paragraphs[0]
    cr1 = cp.add_run("💡 EL BENEFICIO CLAVE PARA EL PROPIETARIO:\n")
    cr1.bold = True
    cr1.font.size = Pt(9.5)
    cr1.font.color.rgb = NAVY
    cr2 = cp.add_run(
        "La emisora mantiene su identidad radial tradicional pero adquiere tecnología de primer nivel nacional, "
        "con una herramienta comercial que permite financiar los costos de operación y generar utilidades mensuales "
        "desde el primer mes de puesta en marcha."
    )
    cr2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. Los 8 Pilares Fundamentales
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Los 8 Pilares Estratégicos de la Plataforma")
    r2.font.size = Pt(14)
    r2.font.color.rgb = NAVY
    r2.font.bold = True

    pilares = [
        ("Pilar 1: Señal de Audio Streaming Inteligente Dual (HD 128 kbps + ECO 64 kbps)",
         "• Modo HD (128 kbps): Sonido estéreo cristalino para hogares y locales comerciales con Wi-Fi.\n"
         "• Modo Ahorro ECO (64 kbps): Específicamente optimizado para botes en altamar y celulares con señal débil 3G/4G sin cortes.\n"
         "• Audio Persistente: El auditor puede recorrer la web, revisar noticias y ver mareas sin que la transmisión se detenga.\n"
         "• Temporizador Sleep Timer: Los auditores pueden programar el apagado automático para quedarse dormidos escuchando la radio."),

        ("Pilar 2: Cadena Radial de Emergencia Comunal en Vivo",
         "• Despliegue de Alerta Prioritaria: Cinta roja con baliza y sirena visual fija en el tope de toda la web.\n"
         "• Emisión en 1 Clic con Plantillas Oficiales: Avisos inmediatos ante Marejadas Anormales, Cortes en Ruta P-72S, Cortes de Suministro Eléctrico y Alertas Meteorológicas SENAPRED.\n"
         "• Sincronización en la Nube: Al emitir una alerta desde el panel, todos los celulares y computadores conectados la reciben en tiempo real."),

        ("Pilar 3: Boletín Marítimo & Mareas Satelitales en Tiempo Real (Quidico -38.243°, -73.492°)",
         "• Conexión Satelital Real: Conectado a los modelos oceanográficos de Open-Meteo & ECMWF/NOAA Marine.\n"
         "• Mediciones en Vivo: Altura del oleaje del Pacífico Sur, velocidad del viento en Nudos (knots), temperatura del agua y fase lunar astronómica para el repunte de mareas.\n"
         "• Semáforo de Puerto Oficial: Estado de Capitanía (Abierto, Precaución, Cerrado) con recomendaciones a pescadores artesanales y recolectoras de orilla."),

        ("Pilar 4: Quidico TV (Canal Audiovisual 1080p)",
         "• Transmisiones en Directo: Espacio preparado para video en vivo (YouTube Live / cámara de cabina / estudio).\n"
         "• Cobertura de Eventos: Festivales locales, campeonatos deportivos, concejos municipales y transmisiones estelares.\n"
         "• Chat Interactivo: Participación y saludos de los vecinos en tiempo real durante los programas."),

        ("Pilar 5: Programación Oficial con Afiches Originales de Alta Resolución",
         "• Los 4 programas oficiales de la estación con sus afiches originales integrados:\n"
         "  1. Las Noticias en Puerto Quidico (Edición Matinal, 06:00 - 09:00 hrs, Lun-Sáb)\n"
         "  2. Las Noticias en Puerto Quidico (Edición Mediodía, 12:00 - 13:00 hrs, Lun-Sáb)\n"
         "  3. Surcando el Lafken (Martes 19:00 - 21:00 hrs, Conduce: Luis Sandoval Antilao)\n"
         "  4. DJ Dino se toma el dial en Quidico (Domingos 10:00 - 14:00 hrs, Dial 91.3 FM)\n"
         "• Indicador Dinámico: La web resalta automáticamente qué programa y locutor está al aire según la hora chilena."),

        ("Pilar 6: Red Comercial & Monetización con Comercializadora Don Nica",
         "• Auspiciador Oficial Activo: Presentación destacada de Comercializadora Don Nica (productos del mar frescos y congelados) con enlaces directos a su WhatsApp de pedidos e Instagram.\n"
         "• Contador de Impacto: Registro de clics y contactos generados para el cliente.\n"
         "• Reportes Automáticos: Botón para generar comprobantes formales en WhatsApp o TXT descargable."),

        ("Pilar 7: Aplicación Móvil PWA (Instalable en Android y iPhone sin tiendas)",
         "• Sin costos de publicación en Google Play ni App Store (ahorro de $100 USD anuales).\n"
         "• Botón 'Instalar App en tu Celular' que añade el icono de la radio en la pantalla de inicio del oyente.\n"
         "• Peso ultra ligero: menos de 2 MB y arranque instantáneo en pantalla completa."),

        ("Pilar 8: Panel de Administración y Control Total para el Dueño (PIN 1051)",
         "• Acceso Protegido: Bloqueo tras 4 intentos fallidos para máxima seguridad contra intrusos.\n"
         "• Fácil acceso: Botón discreto en el encabezado, en el menú móvil o con el atajo de cabina Alt + A.\n"
         "• Móvil y Despachos: Parámetros listos para transmitir desde el celular con Larix Broadcaster o BUTT.\n"
         "• Vúmetro WebRTC: Permite probar la modulación de voz del locutor antes de salir al aire.\n"
         "• Respaldos en 1 Clic: Exportación e importación completa de la configuración en archivo JSON.")
    ]

    for p_head, p_body in pilares:
        h3 = doc.add_heading(level=2)
        r3 = h3.add_run(p_head)
        r3.font.size = Pt(11.5)
        r3.font.color.rgb = BLUE
        r3.font.bold = True
        
        for line in p_body.split('\n'):
            lp = doc.add_paragraph()
            lp.paragraph_format.space_after = Pt(2)
            if line.startswith('•'):
                lp.style = 'List Bullet'
                line = line[1:].strip()
            elif line.startswith('  '):
                lp.style = 'List Bullet 2'
                line = line.strip()
            lp.add_run(line)

    # 4. Modelo Comercial y Tarifario
    doc.add_page_break()
    h3_com = doc.add_heading(level=1)
    r3_com = h3_com.add_run("3. Modelo de Negocio: Proyección de Ingresos Mensuales")
    r3_com.font.size = Pt(14)
    r3_com.font.color.rgb = NAVY
    r3_com.font.bold = True

    doc.add_paragraph(
        "La web está construida con un módulo público de Tarifas Publicitarias que le permite al área comercial "
        "de la radio vender paquetes combinados de mención radial + presencia digital. "
        "A continuación se detalla la estructura comercial implementada:"
    )

    # Commercial Table
    table = doc.add_table(rows=5, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Plan Comercial", "Valor Mensual", "Beneficios Incluidos", "Cupo Sugerido", "Ingreso Proyectado"]
    col_widths = [Inches(1.2), Inches(1.1), Inches(2.5), Inches(1.0), Inches(1.2)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "01244E")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    rows_data = [
        ("Plan Bronce\n(Pyme Local)", "$35.000 / mes", "4 Menciones al aire diarias + Logo y enlace a WhatsApp en la web + Difusión en redes sociales.", "4 clientes", "$140.000"),
        ("Plan Plata\n(Auspicio Bloque)", "$75.000 / mes", "8 Menciones en alta sintonía + Auspicio de Noticias o Boletín Marítimo + Banner web oficial + Reporte mensual.", "3 clientes", "$225.000"),
        ("Plan Oro\n(Cobertura Total)", "$130.000 / mes", "Auspicio estelar rotativo + Banner portada + Mención dominical especial con DJ Dino + Presencia en Quidico TV.", "2 clientes", "$260.000"),
        ("TOTAL ESTIMADO", "—", "9 Auspiciadores locales fidelizados en Tirúa, Cañete y Quidico.", "9 cupos", "$625.000 / mes")
    ]

    for row_idx, rdata in enumerate(rows_data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F1F5F9" if row_idx % 2 == 1 else "FFFFFF"
        if row_idx == 3: # Total row
            bg_color = "E2E8F0"

        for col_idx, text in enumerate(rdata):
            row_cells[col_idx].text = text
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=90, right=90)
            set_cell_border(row_cells[col_idx], 
                            bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'},
                            top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, 1, 3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                if row_idx == 3:
                    run.font.bold = True
                    run.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. Guía de Demostración para el Dueño
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. Protocolo de Demostración en 5 Minutos")
    r4.font.size = Pt(14)
    r4.font.color.rgb = NAVY
    r4.font.bold = True

    steps = [
        ("Paso 1: Abrir la web y probar audio persistente", "Cargar http://localhost:5173/ en un computador o celular. Pulsar 'Play' y mostrar cómo la señal suena con alta fidelidad y nunca se interrumpe al desplazarse por el sitio."),
        ("Paso 2: Mostrar el Boletín Marítimo Satelital", "Hacer clic en la píldora superior 'Puerto Abierto / Pleamar'. Enseñar que los datos de oleaje y viento son reales de Caleta Quidico consultados vía satélite."),
        ("Paso 3: Abrir el Panel de Control Secreto", "Pulsar el botón '🔒 Panel Emisora' en el encabezado o presionar 'Alt + A' en el teclado. Ingresar el PIN 1051."),
        ("Paso 4: El 'Efecto Impacto' de la Alerta de Emergencia", "Ir a la pestaña 'Cadena de Emergencia' y hacer clic en la plantilla 'Marejadas Anormales'. Salir del panel y mostrar cómo la página se transforma en una cadena oficial con sirena y baliza roja. Volver a ingresar y desactivarla en 1 clic."),
        ("Paso 5: Mostrar la sección de Don Nica y descargar el reporte", "Ir a Auspiciadores, enseñar la ficha de Don Nica y pulsar 'Descargar Balance TXT' para que vea el informe contable listo para imprimir.")
    ]

    for s_title, s_desc in steps:
        sp = doc.add_paragraph(style='List Bullet')
        r_st = sp.add_run(s_title + ": ")
        r_st.bold = True
        r_st.font.color.rgb = NAVY
        sp.add_run(s_desc)

    # 6. Ficha Técnica
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. Ficha Técnica y de Infraestructura")
    r5.font.size = Pt(14)
    r5.font.color.rgb = NAVY
    r5.font.bold = True

    tech_items = [
        ("Tecnología Web: ", "React 19 + Tailwind CSS + Lucide Icons + FontAwesome 6 + WebGL Fluid Field."),
        ("Servidor de Streaming: ", "Compatible con Icecast, Shoutcast, SonicPanel y Zeno FM (AAC+ y MP3 directos)."),
        ("Compatibilidad Celular: ", "PWA autónoma para Android, iOS (iPhone/iPad), Smart TV y Windows/Mac."),
        ("Seguridad: ", "Panel protegido por PIN de cabina, encriptación local y rate-limiting contra fuerza bruta."),
        ("Despachos en Terreno: ", "Parámetros integrados para app móvil gratuita Larix Broadcaster (Android/iPhone) y BUTT (PC).")
    ]
    for t_lbl, t_val in tech_items:
        tp = doc.add_paragraph(style='List Bullet')
        rt = tp.add_run(t_lbl)
        rt.bold = True
        rt.font.color.rgb = BLUE
        tp.add_run(t_val)

    # Closing signoff
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_close = doc.add_paragraph()
    r_c1 = p_close.add_run("Radio Puerto Quidico 105.1 FM / 91.3 FM\n")
    r_c1.bold = True
    r_c1.font.color.rgb = NAVY
    r_c2 = p_close.add_run("La voz, la música y las tradiciones de la Costa de Arauco e Isla Mocha al mundo.")
    r_c2.font.italic = True
    r_c2.font.color.rgb = MUTED

    doc.save(output_path)
    print(f"Documento Word guardado exitosamente en: {output_path}")

if __name__ == '__main__':
    target = os.path.abspath('Propuesta_Presentacion_Radio_Puerto_Quidico_105.1FM.docx')
    create_proposal_docx(target)
